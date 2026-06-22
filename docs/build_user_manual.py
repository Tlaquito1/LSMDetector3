from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Manual_de_usuario_Lengua_de_Senas_Mexicana.docx"

FONT = "Aptos"
INK = "2D2523"
SKIN = "A85F55"
PEACH = "FFD9CC"
SAGE = "5E7967"
SAGE_LIGHT = "DCE9DE"
LAVENDER = "78658C"
LAVENDER_LIGHT = "EADFF4"
MUTED = "756763"
LINE = "E6D4CE"


def set_cell_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def set_font(run, size=None, color=INK, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, SKIN, 18, 10),
        "Heading 2": (13, SAGE, 14, 7),
        "Heading 3": (12, LAVENDER, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_running_furniture(section):
    header = section.header.paragraphs[0]
    header.text = "Lengua de Señas Mexicana  |  Manual de usuario"
    set_font(header.runs[0], size=9, color=MUTED, bold=True)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = section.footer.paragraphs[0]
    add_page_field(footer)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("MANUAL DE USUARIO")
    set_font(run, size=11, color=SAGE, bold=True)
    kicker.paragraph_format.space_after = Pt(18)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Lengua de Señas\nMexicana")
    set_font(run, size=30, color=SKIN, bold=True)
    title.paragraph_format.space_after = Pt(10)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Reconocimiento, escritura de frases y entrenamiento desde Android")
    set_font(run, size=14, color=LAVENDER)
    subtitle.paragraph_format.space_after = Pt(28)

    band = doc.add_paragraph()
    set_cell_shading(band, PEACH)
    band.paragraph_format.space_before = Pt(0)
    band.paragraph_format.space_after = Pt(24)
    run = band.add_run("  Guía práctica para usuarios y responsables del proyecto  ")
    set_font(run, size=11, color=INK, bold=True)
    band.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Versión 1.0  |  {date.today().strftime('%d/%m/%Y')}")
    set_font(run, size=10, color=MUTED)

    for _ in range(5):
        doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Documento vivo: se actualizará junto con la aplicación.")
    set_font(run, size=10, color=SAGE, italic=True)
    doc.add_page_break()


def add_callout(doc, label, text, fill=PEACH):
    p = doc.add_paragraph()
    set_cell_shading(p, fill)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"{label}: ")
    set_font(run, bold=True, color=INK)
    run = p.add_run(text)
    set_font(run, color=INK)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def build_manual():
    doc = Document()
    style_document(doc)
    add_running_furniture(doc.sections[0])
    add_cover(doc)

    doc.add_heading("1. Acerca de la aplicación", level=1)
    doc.add_paragraph(
        "Lengua de Señas Mexicana es una aplicación Android que utiliza la cámara, "
        "MediaPipe y muestras de puntos de la mano para reconocer señas, mostrar una "
        "predicción y construir frases en pantalla."
    )
    doc.add_heading("Funciones principales", level=2)
    add_bullets(doc, [
        "Creación de cuenta e inicio de sesión mediante Firebase Authentication.",
        "Detección de 21 puntos de la mano en tiempo real.",
        "Reconocimiento provisional de letras estáticas y señas con movimiento.",
        "Construcción de frases con espacio, borrado y limpieza.",
        "Cambio entre cámara frontal y trasera.",
        "Modo temporal de entrenamiento y exportación de muestras a CSV.",
    ])
    add_callout(
        doc,
        "Estado actual",
        "El reconocimiento utiliza similitud contra muestras guardadas. El entrenamiento "
        "de un modelo final TFLite queda pendiente hasta reunir más datos de distintas personas.",
        fill=LAVENDER_LIGHT,
    )

    doc.add_heading("2. Requisitos y primer inicio", level=1)
    add_bullets(doc, [
        "Dispositivo Android compatible con la versión instalada de la aplicación.",
        "Cámara frontal o trasera disponible.",
        "Conexión a Internet para crear cuenta o iniciar sesión.",
        "Permiso de cámara concedido.",
    ])
    doc.add_heading("Primer inicio", level=2)
    add_steps(doc, [
        "Abre la aplicación Lengua de Señas Mexicana.",
        "Crea una cuenta o inicia sesión con un correo registrado.",
        "Acepta el permiso de cámara cuando Android lo solicite.",
        "Coloca una mano dentro del encuadre y comprueba que aparezcan puntos sobre ella.",
    ])

    doc.add_heading("3. Cuenta y sesión", level=1)
    doc.add_heading("Crear una cuenta", level=2)
    add_steps(doc, [
        "En la pantalla de acceso pulsa Crear cuenta.",
        "Escribe nombre, correo electrónico y una contraseña de al menos seis caracteres.",
        "Confirma la contraseña y pulsa Registrar.",
        "Espera a que Firebase termine el registro; la aplicación abrirá el detector.",
    ])
    doc.add_heading("Iniciar y cerrar sesión", level=2)
    doc.add_paragraph(
        "Firebase conserva la sesión en el dispositivo. La misma cuenta puede usarse en "
        "otros teléfonos con conexión a Internet. Para salir, pulsa Salir en la barra superior."
    )
    add_callout(
        doc,
        "Privacidad",
        "La contraseña es administrada por Firebase y no se guarda en la base SQLite local.",
        fill=SAGE_LIGHT,
    )

    doc.add_heading("4. Usar el detector", level=1)
    add_steps(doc, [
        "Coloca la mano completa frente a la cámara con iluminación uniforme.",
        "Espera a que aparezcan los puntos y líneas sobre las articulaciones.",
        "Mantén una letra estática durante aproximadamente tres segundos para escribirla.",
        "Para una seña dinámica, realiza el movimiento completo y detén la mano brevemente.",
        "Revisa la letra detectada, la confianza aproximada y la barra de confirmación.",
    ])
    doc.add_heading("Controles de frase", level=2)
    add_bullets(doc, [
        "Espacio: agrega una separación entre palabras.",
        "Borrar: elimina el último carácter.",
        "Limpiar: borra toda la frase.",
    ])
    doc.add_heading("Cambiar de cámara", level=2)
    doc.add_paragraph(
        "Pulsa Cámara frontal o Cámara trasera sobre la vista de cámara. Los puntos se "
        "reflejan automáticamente para coincidir con la vista frontal."
    )

    doc.add_heading("5. Entrenar señas temporalmente", level=1)
    doc.add_paragraph(
        "El modo Entrenar permite ampliar las muestras locales mientras se prepara el "
        "modelo final. Las muestras se guardan como coordenadas, no como fotografías."
    )
    doc.add_heading("Letras estáticas", level=2)
    add_steps(doc, [
        "Selecciona una letra de la lista.",
        "Coloca la mano y espera el mensaje Mano detectada.",
        "Pulsa Guardar ejemplo para una muestra o Capturar 100 automáticamente.",
        "Durante la captura automática cambia ligeramente ángulo, distancia y posición.",
    ])
    doc.add_heading("Señas con movimiento", level=2)
    doc.add_paragraph("Las etiquetas dinámicas actuales son J, K, Ñ, Q, X, Z y HOLA.")
    add_steps(doc, [
        "Selecciona una etiqueta dinámica.",
        "Coloca la mano en la posición inicial.",
        "Pulsa Grabar movimiento y ejecuta la trayectoria completa.",
        "Detén la mano y repite desde la posición inicial para crear otro ejemplo.",
    ])
    add_callout(
        doc,
        "Calidad de datos",
        "Conviene recopilar ejemplos de varias personas, velocidades, fondos e iluminaciones. "
        "Muchas capturas casi idénticas no equivalen a datos diversos.",
        fill=PEACH,
    )

    doc.add_heading("6. Datos y exportación", level=1)
    doc.add_paragraph(
        "SQLite guarda las muestras en sign_samples. Cada pose estática contiene 63 valores; "
        "cada movimiento contiene una secuencia de frames separados."
    )
    doc.add_heading("Exportar CSV", level=2)
    add_steps(doc, [
        "Abre Entrenamiento de señas.",
        "Pulsa Exportar CSV.",
        "Conecta el teléfono al equipo y abre Device Explorer en Android Studio.",
        "Busca data/data/com.example.lsmdetector/files/training_exports/sign_samples.csv.",
        "Copia el archivo al proyecto antes de generar una versión distribuible actualizada.",
    ])
    doc.add_paragraph(
        "Las instalaciones nuevas importan las muestras incluidas en assets/sign_samples.csv. "
        "Esto permite que el reconocedor provisional funcione en otro dispositivo."
    )

    doc.add_heading("7. Solución de problemas", level=1)
    doc.add_heading("No aparecen puntos", level=2)
    add_bullets(doc, [
        "Confirma el permiso de cámara en Ajustes de Android.",
        "Usa un fondo con contraste y evita contraluz.",
        "Aleja la mano hasta que se vea completa.",
        "Prueba la otra cámara.",
    ])
    doc.add_heading("Se detectan letras incorrectas", level=2)
    add_bullets(doc, [
        "Mantén estable la pose antes de confirmarla.",
        "Para movimientos, realiza una trayectoria clara y después detén la mano.",
        "Agrega muestras variadas de las clases que se confunden.",
        "Recuerda que la confianza es aproximada mientras no exista el modelo final.",
    ])
    doc.add_heading("No puedo iniciar sesión", level=2)
    add_bullets(doc, [
        "Comprueba la conexión a Internet.",
        "Verifica correo y contraseña.",
        "Si la cuenta era de la versión SQLite anterior, créala nuevamente en Firebase.",
    ])
    doc.add_heading("Android Studio no instala la app", level=2)
    doc.add_paragraph(
        "Activa Depuración USB e instalación por USB en Opciones de desarrollador. En algunos "
        "dispositivos Samsung también debe desactivarse temporalmente el Bloqueador automático."
    )

    doc.add_heading("8. Alcance y recomendaciones", level=1)
    add_bullets(doc, [
        "La aplicación es un prototipo educativo y no sustituye a un intérprete profesional.",
        "El reconocimiento mejora cuando el conjunto incluye múltiples personas.",
        "No compartas archivos de configuración Firebase ni credenciales públicamente.",
        "Conserva respaldos del CSV antes de borrar o reinstalar la aplicación.",
    ])
    add_callout(
        doc,
        "Próxima etapa",
        "Entrenar localmente modelos con los landmarks disponibles, evaluar precisión por clase "
        "e integrar modelos portables en el APK.",
        fill=LAVENDER_LIGHT,
    )

    doc.add_heading("9. Historial del manual", level=1)
    p = doc.add_paragraph()
    run = p.add_run("Versión 1.0 - ")
    set_font(run, bold=True, color=SKIN)
    p.add_run(
        "Documento inicial con cuenta Firebase, detector, frases, cámaras, entrenamiento, "
        "exportación y solución de problemas."
    )

    doc.core_properties.title = "Manual de usuario - Lengua de Señas Mexicana"
    doc.core_properties.subject = "Guía de uso de la aplicación Android"
    doc.core_properties.author = "Proyecto Lengua de Señas Mexicana"
    doc.core_properties.keywords = "LSM, Android, MediaPipe, Firebase, manual"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()
