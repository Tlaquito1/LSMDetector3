from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUTPUT = Path(__file__).with_name("REQUISITOS_FUNCIONALES_Y_NO_FUNCIONALES_LSMDetector_completo.docx")
PDF_OUTPUT = Path(__file__).with_name("REQUISITOS_FUNCIONALES_Y_NO_FUNCIONALES_LSMDetector_completo.pdf")


functional_requirements = [
    "El usuario debe poder crear una cuenta con correo y contraseña para acceder a la aplicación.",
    "El usuario debe poder iniciar sesión con una cuenta registrada y mantener su acceso desde Firebase Authentication.",
    "La aplicación debe mostrar una pantalla inicial con la cámara del dispositivo para iniciar la detección de señas.",
    "El usuario debe poder cambiar entre la cámara frontal y la cámara trasera cuando lo necesite.",
    "La aplicación debe detectar los puntos principales de la mano usando MediaPipe para analizar la posición de los dedos.",
    "La aplicación debe mostrar visualmente los puntos de la mano detectada en la pantalla normal y en la pantalla de entrenamiento.",
    "El usuario debe poder entrenar temporalmente señas del abecedario desde el celular mediante capturas o secuencias de movimiento.",
    "La aplicación debe permitir registrar señas con movimiento, como J, K, Ñ, Q, X, Z, HOLA y otras que requieran desplazamiento.",
    "La aplicación debe reconocer letras estáticas y señas con movimiento para mostrarlas en pantalla como resultado de la detección.",
    "La aplicación debe formar frases en la pantalla inicial cuando una letra o seña se mantenga estable durante el tiempo configurado.",
    "La aplicacion debe importar muestras entrenadas desde un archivo incluido en los assets para que otro celular pueda usar el modelo base.",
    "La aplicacion debe guardar muestras de entrenamiento en una base de datos SQLite local mientras se prepara el modelo final.",
]


non_functional_requirements = [
    "La aplicación debe responder en tiempo casi real para que la detección de señas se sienta fluida durante el uso de la cámara.",
    "La interfaz debe ser intuitiva, clara y visualmente agradable, con colores pastel y tonos piel adecuados para una app educativa.",
    "La aplicación debe evitar detecciones falsas cuando no haya una mano visible en la cámara.",
    "La aplicación debe limpiar los puntos de la mano en pantalla cuando el usuario retire la mano de la cámara.",
    "La detección de señas con movimiento debe tener filtros de estabilidad para no escribir letras antes de completar el movimiento.",
    "La aplicación debe proteger los datos de inicio de sesión usando los servicios seguros de Firebase Authentication.",
    "La aplicación debe ser compatible con dispositivos Android modernos que cuenten con cámara y soporte para CameraX.",
    "El reconocimiento debe funcionar sin depender de conexión a internet para la detección local una vez cargadas las muestras o el modelo.",
    "La base de datos local debe organizar la información de entrenamiento de forma consistente para evitar pérdida o mezcla de muestras.",
    "La aplicación debe mantener buen rendimiento y evitar consumo excesivo de batería mientras usa cámara e inteligencia artificial.",
    "La app debe mostrar mensajes claros cuando falten permisos de cámara o exista algún problema al iniciar la detección.",
    "El proyecto debe poder compilarse desde Android Studio y mantenerse documentado para facilitar futuras modificaciones.",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="DADCE0"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color in (
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_title(doc):
    institution = doc.add_paragraph()
    institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = institution.add_run("TECNOLÓGICO UNIVERSITARIO DE CHALCO")
    run.bold = True
    run.font.size = Pt(14)

    subject = doc.add_paragraph()
    subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subject.add_run("Programación de Móviles - Grupo ISC33")
    run.font.size = Pt(11)

    teacher = doc.add_paragraph()
    teacher.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = teacher.add_run("Docente: Ing. Brenda Tzompantzi Chavez")
    run.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Requisitos Funcionales y No Funcionales de tu Proyecto")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("000000")


def add_intro(doc):
    doc.add_paragraph(
        "Instrucciones: Recuerda la diferencia entre ambos tipos de requisitos (tema 3.7) y "
        "luego completa la tabla de abajo con requisitos reales para la app que estás "
        "desarrollando como proyecto final."
    )
    add_difference_table(doc)


def add_difference_table(doc):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_width(table, [Inches(1.65), Inches(3.0), Inches(1.85)])
    headers = ["Tipo de requisito", "¿Qué describe?", "Ejemplo visto en clase"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_border(cell)
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("000000")

    example_rows = [
        (
            "Funcional",
            'Qué hace la app (las acciones que el usuario puede realizar)',
            '"El usuario debe poder iniciar sesión con Google y subir una foto de perfil."',
        ),
        (
            "No Funcional",
            "Cómo se comporta la app (calidad, rendimiento, compatibilidad)",
            '"La app debe pesar menos de 50 MB y funcionar sin internet."',
        ),
    ]
    for row in example_rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            set_cell_border(cells[idx])
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[idx].paragraphs:
                if idx == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor.from_string("000000")


def add_requirements_table(doc):
    doc.add_paragraph()
    intro = doc.add_paragraph()
    run = intro.add_run("Ahora te toca a ti: escribe al menos 10 requisitos funcionales y 10 no funcionales de tu propio proyecto.")
    run.bold = True

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_width(table, [Inches(0.55), Inches(1.55), Inches(4.4)])

    headers = ["#", "Tipo (Funcional / No Funcional)", "Requisito de TU proyecto"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("000000")

    rows = []
    for requirement in functional_requirements[:10]:
        rows.append(("Funcional", requirement))
    for requirement in non_functional_requirements[:10]:
        rows.append(("No funcional", requirement))

    for idx, (kind, requirement) in enumerate(rows, 1):
        row_cells = table.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = kind
        row_cells[2].text = requirement
        for cell_idx, cell in enumerate(row_cells):
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                if cell_idx in (0, 1):
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_extra_notes(doc):
    return


def add_footer(doc):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Requisitos del proyecto LSMDetector")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("000000")


def build_pdf():
    regular_font = "Helvetica"
    bold_font = "Helvetica-Bold"
    arial = Path("C:/Windows/Fonts/arial.ttf")
    arial_bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if arial.exists() and arial_bold.exists():
        pdfmetrics.registerFont(TTFont("ArialLocal", str(arial)))
        pdfmetrics.registerFont(TTFont("ArialLocal-Bold", str(arial_bold)))
        regular_font = "ArialLocal"
        bold_font = "ArialLocal-Bold"

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=18,
        textColor=colors.black,
        alignment=1,
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        "SubtitleCustom",
        parent=styles["Normal"],
        fontName=regular_font,
        fontSize=11,
        textColor=colors.black,
        alignment=1,
        spaceAfter=14,
    )
    heading_style = ParagraphStyle(
        "HeadingCustom",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=14,
        textColor=colors.black,
        spaceBefore=12,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "BodyCustom",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=9.2,
        leading=11.2,
        spaceAfter=7,
    )
    cell_style = ParagraphStyle(
        "CellCustom",
        parent=styles["BodyText"],
        fontName=regular_font,
        fontSize=8,
        leading=10,
    )
    header_style = ParagraphStyle(
        "HeaderCustom",
        parent=cell_style,
        fontName=bold_font,
        textColor=colors.black,
        alignment=1,
    )

    doc = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
    )

    story = [
        Paragraph("TECNOLÓGICO UNIVERSITARIO DE CHALCO", title_style),
        Paragraph("Programación de Móviles - Grupo ISC33", subtitle_style),
        Paragraph("Docente: Ing. Brenda Tzompantzi Chavez", subtitle_style),
        Paragraph("Requisitos Funcionales y No Funcionales de tu Proyecto", heading_style),
        Paragraph(
            "Instrucciones: Recuerda la diferencia entre ambos tipos de requisitos (tema 3.7) y luego completa "
            "la tabla de abajo con requisitos reales para la app que estás desarrollando como proyecto final.",
            body_style,
        ),
    ]

    difference_rows = [
        [Paragraph("Tipo de requisito", header_style), Paragraph("¿Qué describe?", header_style), Paragraph("Ejemplo visto en clase", header_style)],
        [
            Paragraph("Funcional", cell_style),
            Paragraph("Qué hace la app (las acciones que el usuario puede realizar)", cell_style),
            Paragraph('"El usuario debe poder iniciar sesión con Google y subir una foto de perfil."', cell_style),
        ],
        [
            Paragraph("No Funcional", cell_style),
            Paragraph("Cómo se comporta la app (calidad, rendimiento, compatibilidad)", cell_style),
            Paragraph('"La app debe pesar menos de 50 MB y funcionar sin internet."', cell_style),
        ],
    ]
    difference_table = Table(difference_rows, colWidths=[1.25 * inch, 3.05 * inch, 2.35 * inch], repeatRows=1)
    difference_table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.55, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (0, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(difference_table)
    story.append(Spacer(1, 10))
    story.append(
        Paragraph(
            "Ahora te toca a ti: escribe al menos 10 requisitos funcionales y 10 no funcionales de tu propio proyecto.",
            heading_style,
        )
    )

    rows = [[Paragraph("#", header_style), Paragraph("Tipo", header_style), Paragraph("Requisito de TU proyecto", header_style)]]
    selected_rows = [("Funcional", item) for item in functional_requirements[:10]]
    selected_rows += [("No funcional", item) for item in non_functional_requirements[:10]]
    for idx, (kind, req) in enumerate(selected_rows, 1):
        rows.append([Paragraph(str(idx), cell_style), Paragraph(kind, cell_style), Paragraph(req, cell_style)])

    table = Table(rows, colWidths=[0.35 * inch, 0.95 * inch, 5.35 * inch], repeatRows=1)
    commands = [
        ("GRID", (0, 0), (-1, -1), 0.55, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    table.setStyle(TableStyle(commands))
    story.append(table)
    doc.build(story)


def main():
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_intro(doc)
    add_requirements_table(doc)
    add_extra_notes(doc)
    doc.save(OUTPUT)
    build_pdf()
    print(OUTPUT)
    print(PDF_OUTPUT)


if __name__ == "__main__":
    main()
